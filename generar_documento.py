import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def add_code_block(doc, text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.2)
    set_cell_background(cell, "1E1E1E")
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.left_indent = Inches(0.1)
    p.paragraph_format.right_indent = Inches(0.1)
    
    run = p.add_run(text)
    run.font.name = 'Consolas'
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(220, 220, 220)

doc = docx.Document()

# Margenes
for section in doc.sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

# Estilo de Titulo Principal
title_p = doc.add_paragraph()
title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
title_run = title_p.add_run("Documentación Técnica: Arquitectura y Patrón Singleton")
title_run.font.name = 'Segoe UI'
title_run.font.size = Pt(22)
title_run.font.bold = True
title_run.font.color.rgb = RGBColor(229, 9, 20) # Red acento

sub_p = doc.add_paragraph()
sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub_run = sub_p.add_run("Sistema de Autenticación CineStream (Flask + PostgreSQL)")
sub_run.font.name = 'Segoe UI'
sub_run.font.size = Pt(13)
sub_run.font.color.rgb = RGBColor(100, 100, 100)

doc.add_paragraph().paragraph_format.space_after = Pt(12)

# --- SECCIÓN 1 ---
h1 = doc.add_heading("1. El Patrón de Diseño Singleton", level=1)
h1.style.font.color.rgb = RGBColor(30, 30, 30)

p = doc.add_paragraph()
p.add_run("¿Qué es? ").bold = True
p.add_run("El patrón Singleton es un patrón de diseño creacional que garantiza que una clase tenga únicamente una instancia en todo el ciclo de vida de la aplicación y proporciona un punto de acceso global a ella.")

p = doc.add_paragraph()
p.add_run("¿Dónde está ubicado en el proyecto? ").bold = True
p.add_run("En el archivo ")
p.add_run("database/connection.py").bold = True
p.add_run(" dentro de la clase ")
p.add_run("DatabaseSingleton").bold = True
p.add_run(".")

p = doc.add_paragraph()
p.add_run("¿Por qué y en dónde se usa? ").bold = True
p.add_run("Se utiliza para la conexión a la base de datos PostgreSQL. Abrir múltiples conexiones a una base de datos consume recursos críticos (memoria, sockets de red y tiempo de CPU). Al usar Singleton, toda la aplicación (repositorios, servicios y rutas de Flask) reutiliza exactamente la misma conexión física a PostgreSQL.")

doc.add_paragraph().paragraph_format.space_after = Pt(6)

# Código Singleton
doc.add_heading("Código de database/connection.py Explicado Linea por Línea:", level=2)

add_code_block(doc, """import psycopg2

class DatabaseSingleton:
    _instance = None  # Atributo de clase privado para almacenar la instancia única

    def __new__(cls, *args, **kwargs):
        if not cls._instance:
            # Si la instancia no existe, se crea por primera vez
            cls._instance = super(DatabaseSingleton, cls).__new__(cls)
            cls._instance._connection = psycopg2.connect(
                dbname="login_db",
                user="postgres",
                password="tu_password",
                host="localhost",
                port="5432"
            )
            print("Conexión exitosa a la base de datos (Instancia Singleton creada).")
        return cls._instance

    def get_connection(self):
        return self._connection""")

# Lista explicativa Singleton
doc.add_paragraph()
p = doc.add_paragraph(style='List Bullet')
p.add_run("_instance = None: ").bold = True
p.add_run("Variable de clase que almacena el puntero a la única instancia creada.")

p = doc.add_paragraph(style='List Bullet')
p.add_run("__new__(cls): ").bold = True
p.add_run("Método especial de Python que se ejecuta antes de __init__ para controlar la creación del objeto. Si _instance es None, crea el objeto y establece la conexión con psycopg2; de lo contrario, retorna la instancia previamente creada.")

p = doc.add_paragraph(style='List Bullet')
p.add_run("get_connection(self): ").bold = True
p.add_run("Método público que devuelve la conexión activa hacia PostgreSQL para que la utilicen los repositorios.")

doc.add_paragraph().paragraph_format.space_after = Pt(12)

# --- SECCIÓN 2 ---
doc.add_heading("2. Explicación de las Capas del Proyecto", level=1)

# Repositorio
doc.add_heading("A. Capa de Datos (repositories/user_repository.py)", level=2)
p = doc.add_paragraph()
p.add_run("Esta capa interactúa directamente con PostgreSQL realizando las operaciones SQL CRUD.")

add_code_block(doc, """from database.connection import DatabaseSingleton
from models.user import User

class UserRepository:
    def __init__(self):
        # Consume la instancia única del Singleton
        self.db = DatabaseSingleton().get_connection()

    def create_user(self, user: User) -> bool:
        query = \"\"\"
            INSERT INTO usuarios (nombre, correo, password, rol, activo)
            VALUES (%s, %s, %s, %s, %s);
        \"\"\"
        with self.db.cursor() as cursor:
            cursor.execute(query, (user.nombre, user.correo, user.password, user.rol, user.activo))
            self.db.commit()
            return True

    def find_by_email(self, correo: str) -> User | None:
        query = \"\"\"
            SELECT id, nombre, correo, password, rol, activo
            FROM usuarios WHERE correo = %s;
        \"\"\"
        with self.db.cursor() as cursor:
            cursor.execute(query, (correo,))
            result = cursor.fetchone()
            if result:
                return User(
                    user_id=result[0], nombre=result[1], correo=result[2],
                    password=result[3], rol=result[4], activo=result[5]
                )
            return None""")

# Servicio
doc.add_heading("B. Lógica de Negocio (services/auth_service.py)", level=2)
p = doc.add_paragraph()
p.add_run("Contiene las reglas de negocio, como la verificación de duplicados de correo y el hashing de contraseñas con bcrypt.")

add_code_block(doc, """from repositories.user_repository import UserRepository
from services.password_hasher import PasswordHasher
from models.user import User

class AuthService:
    def __init__(self):
        self.user_repo = UserRepository()
        self.hasher = PasswordHasher()

    def registrar_usuario(self, nombre, correo, password):
        # 1. Validación de duplicados usando el repositorio
        if self.user_repo.find_by_email(correo):
            return False

        # 2. Hashing seguro de la contraseña
        hashed_password = self.hasher.hash_password(password)

        # 3. Creación de entidad y persistencia
        nuevo_usuario = User(
            nombre=nombre, correo=correo,
            password=hashed_password, rol="cliente", activo=True
        )
        return self.user_repo.create_user(nuevo_usuario)""")

# Rutas / App
doc.add_heading("C. Controlador y Rutas (app.py)", level=2)
p = doc.add_paragraph()
p.add_run("Gestiona las peticiones HTTP de Flask, validaciones del lado del servidor (formato de correo y longitud de contraseña) y la renderización de las plantillas HTML.")

add_code_block(doc, """@app.route("/register", methods=["GET", "POST"])
def register():
    if request.method == "POST":
        nombre = request.form.get("nombre", "").strip()
        correo = request.form.get("correo", "").strip()
        password = request.form.get("password", "")
        confirm_password = request.form.get("confirm_password", "")

        # Validaciones de servidor
        if "@" not in correo:
            return render_template("register.html", error="Correo inválido.")
        if len(password) < 6:
            return render_template("register.html", error="Mínimo 6 caracteres.")
        if password != confirm_password:
            return render_template("register.html", error="Las contraseñas no coinciden.")

        if auth_service.registrar_usuario(nombre, correo, password):
            return redirect(url_for("index"))
        return render_template("register.html", error="Correo ya registrado.")

    return render_template("register.html")""")

# Guardar documento
doc.save("Documentacion_CineStream_Singleton.docx")
print("¡Documento 'Documentacion_CineStream_Singleton.docx' creado exitosamente!")